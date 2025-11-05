"""
Quiz to Word Document Converter Service
Converts quiz data to Word document for download
"""
import json
import logging
from pathlib import Path
from typing import Optional, List, Dict, Any
from datetime import datetime
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

logger = logging.getLogger(__name__)


class QuizToWordConverter:
    """
    Convert quiz data (JSON format) to Word document.
    Supports temporary file creation for session-based downloads.
    """
    
    def __init__(self):
        """Initialize the converter"""
        self.temp_folder = Path("session_data/temp")
        self.temp_folder.mkdir(parents=True, exist_ok=True)
    
    def create_word_from_quiz_file(
        self, 
        quiz_json_path: str,
        output_filename: Optional[str] = None
    ) -> str:
        """
        Create Word document from quiz JSON file.
        
        Args:
            quiz_json_path: Path to quiz_data.json file
            output_filename: Optional custom filename for output Word file
            
        Returns:
            Path to generated Word file
            
        Raises:
            FileNotFoundError: If quiz file doesn't exist
            ValueError: If quiz data is invalid
        """
        quiz_path = Path(quiz_json_path)
        
        if not quiz_path.exists():
            raise FileNotFoundError(f"Quiz file not found: {quiz_json_path}")
        
        # Load quiz data
        with open(quiz_path, 'r', encoding='utf-8') as f:
            quiz_data = json.load(f)
        
        # Generate output filename
        if not output_filename:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            output_filename = f"quiz_{timestamp}.docx"
        
        output_path = self.temp_folder / output_filename
        
        # Create Word document
        self._create_word_document(quiz_data, str(output_path))
        
        logger.info(f"Word document created: {output_path}")
        return str(output_path)
    
    def create_word_from_quiz_data(
        self,
        quiz_data: Dict[str, Any],
        output_filename: Optional[str] = None
    ) -> str:
        """
        Create Word document directly from quiz data dict.
        
        Args:
            quiz_data: Quiz data dictionary
            output_filename: Optional custom filename
            
        Returns:
            Path to generated Word file
        """
        if not output_filename:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            output_filename = f"quiz_{timestamp}.docx"
        
        output_path = self.temp_folder / output_filename
        self._create_word_document(quiz_data, str(output_path))
        
        logger.info(f"Word document created: {output_path}")
        return str(output_path)
    
    def _create_word_document(self, quiz_data: Dict[str, Any], output_path: str):
        """
        Internal method to create formatted Word document.
        
        Args:
            quiz_data: Quiz data dictionary
            output_path: Path to save Word file
        """
        doc = Document()
        
        # Set document margins
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)
        
        # Add title
        title = doc.add_heading('BỘ CÂU HỎI KIỂM TRA', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Add metadata if available
        metadata = quiz_data.get('metadata', {})
        if metadata:
            doc.add_paragraph(f"Môn học: {metadata.get('subject', 'N/A')}")
            doc.add_paragraph(f"Chủ đề: {metadata.get('topic', 'N/A')}")
            doc.add_paragraph(f"Số câu hỏi: {metadata.get('total_questions', len(quiz_data.get('questions', [])))}")
            doc.add_paragraph(f"Ngày tạo: {datetime.now().strftime('%d/%m/%Y %H:%M')}")
            doc.add_paragraph()  # Empty line
        
        # Add instructions
        instructions = doc.add_paragraph()
        instructions_run = instructions.add_run('HƯỚNG DẪN: ')
        instructions_run.bold = True
        instructions.add_run('Vui lòng đọc kỹ từng câu hỏi và chọn đáp án đúng nhất.')
        doc.add_paragraph()  # Empty line
        
        # Add questions
        questions = quiz_data.get('questions', [])
        
        for idx, question in enumerate(questions, 1):
            # Question number and text
            question_para = doc.add_paragraph()
            question_num = question_para.add_run(f'Câu {idx}: ')
            question_num.bold = True
            question_num.font.size = Pt(12)
            
            question_text = question.get('title', question.get('question', ''))
            question_para.add_run(question_text)
            question_para.style = 'List Number'
            
            # Question type indicator
            question_type = question.get('type', 'multiple_choice')
            if question_type == 'essay':
                type_para = doc.add_paragraph(f'(Câu hỏi tự luận)')
                type_para.style = 'List Bullet'
                doc.add_paragraph()  # Empty line for answer
                doc.add_paragraph()
                doc.add_paragraph()
            else:
                # Multiple choice options
                options = question.get('options', [])
                option_labels = ['A', 'B', 'C', 'D', 'E', 'F']
                
                for opt_idx, option in enumerate(options):
                    if opt_idx < len(option_labels):
                        option_para = doc.add_paragraph()
                        option_label = option_para.add_run(f'{option_labels[opt_idx]}. ')
                        option_label.bold = True
                        option_para.add_run(str(option))
                        option_para.style = 'List Bullet'
                
                # Correct answer (optional - for teacher's copy)
                if 'answer' in question:
                    answer_para = doc.add_paragraph()
                    answer_run = answer_para.add_run(f'Đáp án đúng: {question["answer"]}')
                    answer_run.font.color.rgb = RGBColor(0, 128, 0)  # Green color
                    answer_run.bold = True
                    answer_run.font.size = Pt(10)
                
                # Answer explanation (optional)
                if 'answer_explanation' in question:
                    explain_para = doc.add_paragraph()
                    explain_run = explain_para.add_run(f'Giải thích: {question["answer_explanation"]}')
                    explain_run.font.color.rgb = RGBColor(0, 0, 255)  # Blue color
                    explain_run.font.size = Pt(10)
                    explain_run.italic = True
            
            doc.add_paragraph()  # Empty line between questions
        
        # Add footer
        doc.add_paragraph()
        footer = doc.add_paragraph('--- HẾT ---')
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer_run = footer.runs[0]
        footer_run.bold = True
        
        # Save document
        doc.save(output_path)
        logger.info(f"Word document saved to: {output_path}")
    
    def cleanup_old_files(self, max_age_hours: int = 24):
        """
        Clean up old Word files in temp folder.
        
        Args:
            max_age_hours: Maximum age of files to keep (default 24 hours)
        """
        try:
            current_time = datetime.now()
            for file_path in self.temp_folder.glob("quiz_*.docx"):
                file_age = current_time - datetime.fromtimestamp(file_path.stat().st_mtime)
                if file_age.total_seconds() > (max_age_hours * 3600):
                    file_path.unlink()
                    logger.info(f"Cleaned up old file: {file_path}")
        except Exception as e:
            logger.error(f"Error cleaning up files: {e}")


# Convenience function
def convert_quiz_to_word(
    quiz_json_path: Optional[str] = None,
    quiz_data: Optional[Dict[str, Any]] = None,
    output_filename: Optional[str] = None
) -> str:
    """
    Convenience function to convert quiz to Word.
    
    Args:
        quiz_json_path: Path to quiz JSON file (if loading from file)
        quiz_data: Quiz data dict (if already loaded)
        output_filename: Optional custom filename
        
    Returns:
        Path to generated Word file
    """
    converter = QuizToWordConverter()
    
    if quiz_json_path:
        return converter.create_word_from_quiz_file(quiz_json_path, output_filename)
    elif quiz_data:
        return converter.create_word_from_quiz_data(quiz_data, output_filename)
    else:
        raise ValueError("Must provide either quiz_json_path or quiz_data")
